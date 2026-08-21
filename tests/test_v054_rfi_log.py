from __future__ import annotations

import copy
import os
import zipfile
from datetime import date, time, timedelta

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.datavalidation import DataValidation

from docflow.clipboard import preview_plan_table
from docflow.fs import apply as fsapply
from docflow.ops import rfi_log_add
from docflow.ops.rfi_log_add import (
    AddRfiLogRows,
    WorkbookLockedError,
    _assert_not_locked,
    _clean_workbook_text,
    _default_workbook,
    _read_source,
)


CIV_522_DESCRIPTION = """Before concrete pouring check that:
•\tDisposition, dimension, number and diameter of reinforcing steel, bending, embedded elements, anchor bolts and spacers for reinforcement cover installed according to drawings (IFC)
•\tGrounding system installed.
•\tCable duct installed.
•\tConstruction joints installed.
•\tElevation control performed"""

CIV_523_DESCRIPTION = """During concrete pouring check:
•\tPouring method, vibration, continuity, volumetric efficiency
•\tSlump test, air temperature, control of concrete temperature in situ, sampling, and delivery notes for required mix design.
•\tCheck concrete air content,
•\tTake cube samples for strength test
•\tConcrete traceability control activities"""


def test_default_workbook_finds_onedrive_relocated_systemu(tmp_path, monkeypatch):
    home = tmp_path / "user"
    workbook = (
        home
        / "OneDrive - AZFEN LLC"
        / "Desktop"
        / "Rohtolos"
        / "SystemU"
        / "RFI Log.xlsx"
    )
    workbook.parent.mkdir(parents=True)
    workbook.touch()
    monkeypatch.delenv("DOCFLOW_RFI_LOG", raising=False)
    monkeypatch.setattr(rfi_log_add.pathlib.Path, "home", lambda: home)

    assert _default_workbook() == str(workbook)


def _make_workbook(path):
    wb = Workbook()
    civ = wb.active
    civ.title = "CIV"
    mec = wb.create_sheet("MEC")
    lookup = wb.create_sheet("Sheet1")

    civ_headers = [
        "RFI No.",
        "ITP No:",
        "Item No.",
        "INSPECTION DESCRIPTION",
        "FORMS",
        "TAG NO",
        "AREA",
        "DATE OF ISSUE",
        "INSPECTION DATE",
        "EXPIRY DATE",
        "INSPECTION START TIME",
        "INSP. TYPE AZFEN",
        "INSP. TYPE INTEC",
        "Inspection Location",
        "SUBCONTRACTOR",
        "Inspection Status",
        "RFI VALIDITY",
    ]
    mec_headers = civ_headers[:7] + ["ROW No."] + civ_headers[7:]
    for column, value in enumerate(civ_headers, 1):
        civ.cell(9, column).value = value
    for column, value in enumerate(mec_headers, 1):
        mec.cell(9, column).value = value

    inspection_day = date.today() + timedelta(days=1)
    civ_values = [
        "AZE-SUN-INT-QMS-RFI-CIV-0007",
        "AZE-SUN-INT-QMS-ITP-CIV-0005",
        "CIV-5.2.1",
        "Civil description",
        "CON-CIV-001",
        "N12",
        "North Site",
        date.today(),
        inspection_day,
        date.today() + timedelta(days=8),
        time(17, 30),
        "H",
        "W",
        "N3",
        "AZFEN",
        "Ongoing",
        "Valid",
    ]
    for column, value in enumerate(civ_values, 1):
        civ.cell(10, column).value = value

    mec_values = [
        "AZE-SUN-INT-QMS-RFI-MEC-0020",
        "AZE-SUN-INT-QMS-ITP-MEC-0003",
        "MEC-3.2.1",
        "Mechanical description",
        "CON-MEC-002",
        "PCS-01",
        "S1",
        "",
        date.today(),
        inspection_day,
        date.today() + timedelta(days=8),
        time(9, 0),
        "H",
        "S",
        "N3",
        "AZFEN",
        "Ongoing",
        "Valid",
    ]
    for column, value in enumerate(mec_values, 1):
        mec.cell(10, column).value = value

    # Ayrilmis tablo kapasitesi ve "dokunma" kolonlari.
    civ["J11"] = "KEEP EXPIRY"
    civ["P11"] = "KEEP STATUS"
    civ["Q11"] = "KEEP VALIDITY"
    civ["Q30"] = ""
    mec["K11"] = "KEEP EXPIRY"
    mec["Q11"] = "KEEP STATUS"
    mec["R11"] = "KEEP VALIDITY"
    mec["R30"] = ""

    lookup["K1"] = "Item"
    lookup["L1"] = "Description"
    lookup["M1"] = "Forms"
    lookup["N1"] = "Item"
    lookup["O1"] = "AZFEN"
    lookup["P1"] = "Item"
    lookup["Q1"] = "INTEC"
    lookup["B1"] = "AZE˗SUN˗INT˗QMS˗ITP˗CIV˗0005"
    lookup["C1"] = "AZE˗SUN˗INT˗QMS˗ITP˗MEC˗0003"
    civ["D11"] = '=IFERROR(VLOOKUP(C11,Sheet1!K:L,2,0),"")'
    civ["E11"] = '=IFERROR(VLOOKUP(D11,Sheet1!L:M,2,0),"")'
    civ["L11"] = '=IFERROR(VLOOKUP(C11,Sheet1!N:O,2,0),"")'
    civ["M11"] = '=IFERROR(VLOOKUP(C11,Sheet1!P:Q,2,0),"")'
    lookup.append([None] * 10 + [
        "CIV-5.2.1",
        "Civil description",
        "CON-CIV-001",
        "CIV-5.2.1",
        "H",
        "CIV-5.2.1",
        "W",
    ])
    lookup.append([None] * 10 + [
        "MEC-3.2.1",
        "Mechanical description",
        "CON-MEC-002",
        "MEC-3.2.1",
        "H",
        "MEC-3.2.1",
        "S",
    ])
    lookup.append([None] * 10 + [
        "CIV-5.2.2",
        CIV_522_DESCRIPTION,
        "CON-CIV-001, CON-CIV-009, QM-FRM-027",
        "CIV-5.2.2",
        "H",
        "CIV-5.2.2",
        "W/R",
    ])
    lookup.append([None] * 10 + [
        "CIV-5.2.3",
        CIV_523_DESCRIPTION,
        "CON-CIV-010, Concrete Delivery Tickets",
        "CIV-5.2.3",
        "H",
        "CIV-5.2.3",
        "W/R",
    ])
    wb.save(path)
    wb.close()


def test_civ_522_523_presets_are_complete_and_repair_mojibake(tmp_path):
    path = tmp_path / "RFI Log.xlsx"
    _make_workbook(path)

    source = _read_source(path, "CIV")
    expected = {
        "CIV-5.2.2": {
            "description": CIV_522_DESCRIPTION,
            "forms": "CON-CIV-001, CON-CIV-009, QM-FRM-027",
            "azfenType": "H",
            "intecType": "W/R",
        },
        "CIV-5.2.3": {
            "description": CIV_523_DESCRIPTION,
            "forms": "CON-CIV-010, Concrete Delivery Tickets",
            "azfenType": "H",
            "intecType": "W/R",
        },
    }
    assert {item: source["presets"][item] for item in expected} == expected

    mojibake = CIV_522_DESCRIPTION.encode("utf-8").decode("cp1252")
    assert "â€¢" in mojibake
    assert _clean_workbook_text(mojibake) == CIV_522_DESCRIPTION
    assert _clean_workbook_text(CIV_522_DESCRIPTION) == CIV_522_DESCRIPTION

    operation = AddRfiLogRows()
    plan = operation.preview({
        "workbook": str(path),
        "sheet_counts": {"CIV": 1},
        "overflow_next_day": True,
    })
    auto_fill = plan.meta["editor"]["autoFill"]["values"]
    assert {item: auto_fill[item] for item in expected} == expected

    edited = copy.deepcopy(plan.meta["editor"]["rows"])
    edited[0].update(expected["CIV-5.2.2"])
    edited[0]["itemNo"] = "CIV-5.2.2"
    edited[0]["description"] = mojibake
    revised = operation.revise(plan, {"rows": edited})
    revised_row = revised.meta["editor"]["rows"][0]
    assert revised_row["description"] == CIV_522_DESCRIPTION
    assert "â€¢" not in revised_row["description"]


def test_editor_supports_bulk_selection_hidden_auto_fields_and_pdl(tmp_path):
    path = tmp_path / "RFI Log.xlsx"
    _make_workbook(path)
    operation = AddRfiLogRows()
    plan = operation.preview(
        {
            "workbook": str(path),
            "sheet_counts": {"CIV": 2},
            "overflow_next_day": True,
        }
    )

    editor = plan.meta["editor"]
    assert editor["selectionEnabled"] is True
    assert editor["bulkEnabled"] is True
    assert editor["selectionGroupFields"] == ["_sheet", "rfiNo"]
    assert editor["detailsEnabled"] is True
    assert editor["detailsDefaultVisible"] is False
    detail_fields = {
        column["key"]
        for column in editor["columns"]
        if column.get("detail") is True
    }
    assert detail_fields == {"description", "forms", "azfenType", "intecType"}
    assert editor["autocomplete"]["subcontractor"][:2] == ["AZFEN", "PDL"]

    rows = copy.deepcopy(editor["rows"])
    for row in rows:
        row["subcontractor"] = "PDL"
    revised = operation.revise(plan, {"rows": rows})
    assert {
        row["subcontractor"] for row in revised.meta["editor"]["rows"]
    } == {"PDL"}


@pytest.mark.parametrize(
    ("defect", "message"),
    [
        ("missing", "esleme verisi eksik"),
        ("mismatch", "Item eslemesi tutarsiz"),
        ("conflict", "cakisan Item eslemesi"),
    ],
)
def test_lookup_validation_rejects_bad_data(tmp_path, defect, message):
    path = tmp_path / f"RFI Log-{defect}.xlsx"
    _make_workbook(path)
    workbook = load_workbook(path)
    try:
        lookup = workbook["Sheet1"]
        if defect == "missing":
            lookup["M4"] = ""
        elif defect == "mismatch":
            lookup["N4"] = "CIV-5.2.3"
        else:
            lookup.append([None] * 10 + [
                "CIV-5.2.2",
                "Conflicting description",
                "CON-CIV-001",
                "CIV-5.2.2",
                "H",
                "CIV-5.2.2",
                "W/R",
            ])
        workbook.save(path)
    finally:
        workbook.close()

    with pytest.raises(ValueError, match=message):
        _read_source(path, "CIV")


def _add_ele_prc_sheets(path):
    workbook = load_workbook(path)
    try:
        lookup = workbook["Sheet1"]
        headers = [
            "RFI No.",
            "ITP No:",
            "Item No.",
            "INSPECTION DESCRIPTION",
            "FORMS",
            "TAG NO",
            "AREA",
            "DATE OF ISSUE",
            "INSPECTION DATE",
            "EXPIRY DATE",
            "INSPECTION START TIME",
            "INSP. TYPE AZFEN",
            "INSP. TYPE INTEC",
            "INSP. LOC./AREA",
            "SUBCONTRACTOR",
            "Inspection Status",
            "RFI VALIDITY",
        ]
        existing = {
            "ELE": (
                25,
                "AZE-SUN-INT-QMS-ITP-ELE-0005",
                "ELE-5.1.6",
                "Electrical description",
                "CON-ELE-001",
                "R",
                "W",
            ),
            "PRC": (
                71,
                "AZE-SUN-INT-QMS-ITP-PRC-0002",
                "PRC-2.2.3",
                "Procurement description",
                "CON-MT-001",
                "H",
                "W",
            ),
        }
        for discipline, (
            number,
            itp,
            item,
            description,
            forms,
            azfen_type,
            intec_type,
        ) in existing.items():
            sheet = workbook.create_sheet(discipline)
            for column, value in enumerate(headers, 1):
                sheet.cell(9, column).value = value
            values = [
                f"AZE-SUN-INT-QMS-RFI-{discipline}-{number:04d}",
                itp,
                item,
                description,
                forms,
                "N1",
                "North Site",
                date.today() - timedelta(days=2),
                date.today() - timedelta(days=1),
                date.today(),
                time(10, 0),
                azfen_type,
                intec_type,
                "Laydown",
                "AZFEN",
                "Ongoing",
                "Valid",
            ]
            for column, value in enumerate(values, 1):
                sheet.cell(10, column).value = value
            sheet["D11"] = '=IFERROR(VLOOKUP(C11,Sheet1!K:L,2,0),"")'
            sheet["E11"] = '=IFERROR(VLOOKUP(D11,Sheet1!L:M,2,0),"")'
            sheet["L11"] = '=IFERROR(VLOOKUP(C11,Sheet1!N:O,2,0),"")'
            sheet["M11"] = '=IFERROR(VLOOKUP(C11,Sheet1!P:Q,2,0),"")'
            sheet["J11"] = f"KEEP {discipline} EXPIRY"
            sheet["P11"] = f"KEEP {discipline} STATUS"
            sheet["Q11"] = f"KEEP {discipline} VALIDITY"
            sheet["Q30"] = ""

        lookup["D1"] = "AZE-SUN-INT-QMS-ITP-ELE-0005"
        lookup["E1"] = "AZE-SUN-INT-QMS-ITP-PRC-0002"
        lookup.append(
            [None] * 10
            + [
                "ELE-5.1.6",
                "Electrical description",
                "CON-ELE-001",
                "ELE-5.1.6",
                "R",
                "ELE-5.1.6",
                "W",
            ]
        )
        lookup.append(
            [None] * 10
            + [
                "PRC-2.2.3",
                "Procurement description",
                "CON-MT-001",
                "PRC-2.2.3",
                "H",
                "PRC-2.2.3",
                "W",
            ]
        )

        for sheet_name in ("CIV", "MEC", "ELE", "PRC"):
            sheet = workbook[sheet_name]
            custom = DataValidation(
                type="custom",
                formula1="COUNTIF($A$10:$A$30,A10)<=1",
            )
            item_list = DataValidation(type="list", formula1="INDIRECT(B10)")
            sheet.add_data_validation(custom)
            sheet.add_data_validation(item_list)
            custom.add("A10:A30")
            item_list.add("C10:C30")
        workbook.save(path)
    finally:
        workbook.close()


def test_preview_numbers_schedule_and_overflow(tmp_path):
    path = tmp_path / "RFI Log.xlsx"
    _make_workbook(path)
    plan = AddRfiLogRows().preview({
        "workbook": str(path),
        "discipline": "CIV",
        "count": 2,
        "overflow_next_day": True,
    })

    rows = plan.meta["editor"]["rows"]
    assert [row["rfiNo"] for row in rows] == [
        "AZE-SUN-INT-QMS-RFI-CIV-0008",
        "AZE-SUN-INT-QMS-RFI-CIV-0009",
    ]
    assert rows[0]["inspectionDate"] == (
        date.today() + timedelta(days=2)
    ).strftime("%d-%b-%y")
    assert [row["startTime"] for row in rows] == ["09:00", "09:30"]
    assert any(w.key == "warn.rfi.overflowMoved" for w in plan.warnings)
    assert plan.meta["editor"]["layout"] == "compact"
    columns = {
        column["key"]: column for column in plan.meta["editor"]["columns"]
    }
    assert columns["itpNo"]["display"] == "shortNumber"
    assert columns["itemNo"]["display"] == "itemSuffix"
    assert columns["dateOfIssue"]["kind"] == "compactDate"
    assert columns["inspectionDate"]["kind"] == "compactDate"
    assert rows[0]["itpNo"] == "AZE˗SUN˗INT˗QMS˗ITP˗CIV˗0005"
    assert plan.meta["editor"]["copySource"]["sheet"] == "CIV"
    assert plan.meta["editor"]["copySource"]["columns"] == [
        1, 2, 3, 4, 5, 6, 7, 8, 9, 11, 12, 13, 14, 15
    ]


def test_composite_tag_uses_leading_site_marker():
    assert rfi_log_add._infer_area("N3_PCS-02_S7/S8") == "North Site"
    assert rfi_log_add._infer_area("S4_PCS-01_N1") == "South Site"


def test_shifted_lookup_block_is_discovered_and_formulas_are_repaired(
    tmp_path,
    monkeypatch,
):
    path = tmp_path / "RFI Log.xlsx"
    _make_workbook(path)
    workbook = load_workbook(path)
    try:
        # Production gained an extra ITP list column, moving K:Q to L:R.
        workbook["Sheet1"].insert_cols(11)
        workbook.save(path)
    finally:
        workbook.close()

    monkeypatch.setattr(fsapply, "JOURNAL_DIR", tmp_path / "journal")
    operation = AddRfiLogRows()
    plan = operation.preview({
        "workbook": str(path),
        "discipline": "CIV",
        "count": 1,
        "overflow_next_day": True,
    })

    assert plan.steps[0].meta["lookupColumns"]["item"] == 12
    result = operation.apply(plan)
    assert result.failed == 0

    formulas = load_workbook(path, read_only=True, data_only=False)
    try:
        civ = formulas["CIV"]
        assert civ["D11"].value == '=IFERROR(VLOOKUP(C11,Sheet1!L:M,2,0),"")'
        assert civ["E11"].value == '=IFERROR(VLOOKUP(D11,Sheet1!M:N,2,0),"")'
        assert civ["L11"].value == '=IFERROR(VLOOKUP(C11,Sheet1!O:P,2,0),"")'
        assert civ["M11"].value == '=IFERROR(VLOOKUP(C11,Sheet1!Q:R,2,0),"")'
    finally:
        formulas.close()

    calculated = load_workbook(path, read_only=True, data_only=True)
    try:
        civ = calculated["CIV"]
        assert civ["D11"].value == "Civil description"
        assert civ["E11"].value == "CON-CIV-001"
        assert civ["L11"].value == "H"
        assert civ["M11"].value == "W"
    finally:
        calculated.close()


def test_email_prefills_validated_rows_and_can_duplicate_one_rfi(tmp_path):
    path = tmp_path / "RFI Log.xlsx"
    _make_workbook(path)
    plan = AddRfiLogRows().preview({
        "workbook": str(path),
        "discipline": "MEC",
        "count": 3,
        "email_provider": "offline",
        "email_text": (
            "Create 1 CIV RFI number\n"
            "ITP No: 5\n"
            "Item: CIV-5.2.1\n"
            "Tag No: N12, S14\n"
            "Inspection Location: Foundation A"
        ),
        "overflow_next_day": True,
    })

    rows = plan.meta["editor"]["rows"]
    assert len(rows) == 2
    assert rows[0]["rfiNo"] == rows[1]["rfiNo"]
    assert [row["tagNo"] for row in rows] == ["N12", "S14"]
    assert [row["area"] for row in rows] == ["North Site", "South Site"]
    assert all(row["itemNo"] == "CIV-5.2.1" for row in rows)
    assert all(row["location"] == "Foundation A" for row in rows)
    assert plan.meta["editor"]["copySource"]["rfiNumbers"] == [
        rows[0]["rfiNo"],
        rows[1]["rfiNo"],
    ]
    assert any(w.key == "warn.rfi.emailExtracted" for w in plan.warnings)


def test_revise_repairs_mojibake_itp_separator(tmp_path):
    path = tmp_path / "RFI Log.xlsx"
    _make_workbook(path)
    operation = AddRfiLogRows()
    plan = operation.preview({
        "workbook": str(path),
        "discipline": "CIV",
        "count": 1,
        "overflow_next_day": True,
    })
    rows = copy.deepcopy(plan.meta["editor"]["rows"])
    expected = rows[0]["itpNo"]
    rows[0]["itpNo"] = expected.encode("utf-8").decode("cp1252")

    revised = operation.revise(plan, {"rows": rows})

    assert revised.meta["editor"]["rows"][0]["itpNo"] == expected
    assert revised.steps[0].meta["records"][0]["itpNo"] == expected


def test_apply_duplicate_preserves_untouched_columns_and_undo(tmp_path, monkeypatch):
    path = tmp_path / "RFI Log.xlsx"
    _make_workbook(path)
    monkeypatch.setattr(fsapply, "JOURNAL_DIR", tmp_path / "journal")
    original = fsapply.file_sha256(path)
    operation = AddRfiLogRows()
    plan = operation.preview({
        "workbook": str(path),
        "discipline": "CIV",
        "count": 1,
        "overflow_next_day": True,
    })
    rows = copy.deepcopy(plan.meta["editor"]["rows"])
    duplicate = dict(rows[0])
    duplicate["_rowId"] = "duplicate"
    duplicate["area"] = "South Site"
    rows.append(duplicate)
    revised = operation.revise(plan, {"rows": rows})
    result = operation.apply(revised)

    assert result.failed == 0
    assert result.applied == 2
    journal = fsapply.get_journal(result.journal_id)
    backup = journal["steps"][0]["src"]
    with zipfile.ZipFile(backup) as before, zipfile.ZipFile(path) as after:
        assert before.namelist() == after.namelist()
        for name in before.namelist():
            if name != "xl/worksheets/sheet1.xml":
                assert before.read(name) == after.read(name), name

    wb = load_workbook(path, data_only=False)
    try:
        civ = wb["CIV"]
        assert civ["A11"].value == "AZE-SUN-INT-QMS-RFI-CIV-0008"
        assert civ["A12"].value == "AZE-SUN-INT-QMS-RFI-CIV-0008"
        assert civ["G11"].value in (None, "")
        assert civ["G12"].value == "South Site"
        assert civ["D11"].value == '=IFERROR(VLOOKUP(C11,Sheet1!K:L,2,0),"")'
        assert civ["J11"].value == "KEEP EXPIRY"
        assert civ["P11"].value == "KEEP STATUS"
        assert civ["Q11"].value == "KEEP VALIDITY"
    finally:
        wb.close()

    calculated = load_workbook(path, read_only=True, data_only=True)
    try:
        civ = calculated["CIV"]
        assert civ["D11"].value == "Civil description"
        assert civ["E11"].value == "CON-CIV-001"
        assert civ["L11"].value == "H"
        assert civ["M11"].value == "W"
    finally:
        calculated.close()

    undone = fsapply.undo(result.journal_id)
    assert undone.failed == 0
    assert fsapply.file_sha256(path) == original


def test_mec_writes_after_row_number_without_touching_it(tmp_path, monkeypatch):
    path = tmp_path / "RFI Log.xlsx"
    _make_workbook(path)
    monkeypatch.setattr(fsapply, "JOURNAL_DIR", tmp_path / "journal")
    operation = AddRfiLogRows()
    plan = operation.preview({
        "workbook": str(path),
        "discipline": "MEC",
        "count": 1,
        "overflow_next_day": True,
    })
    result = operation.apply(plan)
    assert result.failed == 0

    wb = load_workbook(path, data_only=False)
    try:
        mec = wb["MEC"]
        assert mec["A11"].value == "AZE-SUN-INT-QMS-RFI-MEC-0021"
        assert mec["H11"].value is None
        assert mec["I11"].value.date() == date.today()
        assert mec["P11"].value == "AZFEN"
        assert mec["K11"].value == "KEEP EXPIRY"
        assert mec["Q11"].value == "KEEP STATUS"
        assert mec["R11"].value == "KEEP VALIDITY"
    finally:
        wb.close()


def test_apply_reports_workbook_lock_without_exposing_os_error(
    tmp_path,
    monkeypatch,
):
    path = tmp_path / "RFI Log.xlsx"
    _make_workbook(path)
    operation = AddRfiLogRows()
    plan = operation.preview({
        "workbook": str(path),
        "discipline": "CIV",
        "count": 1,
        "overflow_next_day": True,
    })
    original = fsapply.file_sha256(path)

    def locked(_path):
        raise WorkbookLockedError("sensitive operating system error")

    monkeypatch.setattr(rfi_log_add, "_assert_not_locked", locked)
    result = operation.apply(plan)

    assert result.failed == 1
    assert result.applied == 0
    assert result.errors[0].key == "error.rfi.workbookLocked"
    assert result.errors[0].subject == path.name
    assert "sensitive" not in result.errors[0].subject
    assert fsapply.file_sha256(path) == original

    access_warning = rfi_log_add._apply_error_warning(
        PermissionError("password or operating system detail"),
        path,
    )
    assert access_warning.key == "error.rfi.workbookLocked"
    assert access_warning.subject == path.name


def test_windows_lock_check_accepts_a_free_workbook(tmp_path):
    if os.name != "nt":
        return
    path = tmp_path / "free.xlsx"
    path.write_bytes(b"free")

    _assert_not_locked(path)


def test_multi_sheet_preview_apply_copy_sources_and_undo(tmp_path, monkeypatch):
    path = tmp_path / "RFI Log.xlsx"
    _make_workbook(path)
    _add_ele_prc_sheets(path)
    monkeypatch.setattr(fsapply, "JOURNAL_DIR", tmp_path / "journal")
    original = fsapply.file_sha256(path)

    operation = AddRfiLogRows()
    plan = operation.preview(
        {
            "workbook": str(path),
            "sheet_counts": {"CIV": 1, "MEC": 1, "ELE": 2, "PRC": 1},
            "overflow_next_day": True,
        }
    )

    rows = plan.meta["editor"]["rows"]
    assert [row["_sheet"] for row in rows] == [
        "CIV",
        "MEC",
        "ELE",
        "ELE",
        "PRC",
    ]
    assert [row["rfiNo"] for row in rows] == [
        "AZE-SUN-INT-QMS-RFI-CIV-0008",
        "AZE-SUN-INT-QMS-RFI-MEC-0021",
        "AZE-SUN-INT-QMS-RFI-ELE-0026",
        "AZE-SUN-INT-QMS-RFI-ELE-0027",
        "AZE-SUN-INT-QMS-RFI-PRC-0072",
    ]
    # Each sheet continues only its own schedule. CIV wraps after its existing
    # 17:30 entry, MEC continues after 09:00, and ELE/PRC start independently.
    assert [(row["inspectionDate"], row["startTime"]) for row in rows] == [
        ((date.today() + timedelta(days=2)).strftime("%d-%b-%y"), "09:00"),
        ((date.today() + timedelta(days=1)).strftime("%d-%b-%y"), "09:30"),
        ((date.today() + timedelta(days=1)).strftime("%d-%b-%y"), "09:00"),
        ((date.today() + timedelta(days=1)).strftime("%d-%b-%y"), "09:30"),
        ((date.today() + timedelta(days=1)).strftime("%d-%b-%y"), "09:00"),
    ]
    assert any(
        warning.key == "warn.rfi.overflowMoved" and warning.subject == "1"
        for warning in plan.warnings
    )
    assert set(plan.meta["editor"]["copySources"]) == {
        "CIV",
        "MEC",
        "ELE",
        "PRC",
    }
    mec_source = plan.meta["editor"]["copySources"]["MEC"]
    assert mec_source["columns"] == [
        1, 2, 3, 4, 5, 6, 7, 8, 9, 10, 12, 13, 14, 15, 16
    ]
    assert mec_source["fields"][7] == "rowNo"
    assert mec_source["headers"][7] == "ROW No."
    assert "rowNo" not in plan.meta["editor"]["copySources"]["CIV"]["fields"]
    assert len(plan.steps) == 1
    assert len(plan.steps[0].meta["batches"]) == 4

    edited = copy.deepcopy(rows)
    edited[1]["rowNo"] = "MEC-ROW-01"
    duplicate = dict(edited[-1])
    duplicate["_rowId"] = "PRC-duplicate"
    duplicate["area"] = "South Site"
    edited.append(duplicate)
    revised = operation.revise(plan, {"rows": edited})
    result = operation.apply(revised)
    assert result.failed == 0
    assert result.applied == 6

    workbook = load_workbook(path, data_only=False)
    try:
        expected = {
            "CIV": (11, "AZE-SUN-INT-QMS-RFI-CIV-0008"),
            "MEC": (11, "AZE-SUN-INT-QMS-RFI-MEC-0021"),
            "ELE": (11, "AZE-SUN-INT-QMS-RFI-ELE-0026"),
            "PRC": (11, "AZE-SUN-INT-QMS-RFI-PRC-0072"),
        }
        for sheet_name, (row_number, rfi) in expected.items():
            sheet = workbook[sheet_name]
            assert sheet.cell(row_number, 1).value == rfi
            assert "VLOOKUP" in str(sheet.cell(row_number, 4).value)
            assert "VLOOKUP" in str(
                sheet.cell(
                    row_number,
                    rfi_log_add.LAYOUTS[sheet_name].azfen_col,
                ).value
            )
        assert workbook["MEC"]["H11"].value == "MEC-ROW-01"
        assert workbook["ELE"]["J11"].value == "KEEP ELE EXPIRY"
        assert workbook["PRC"]["J11"].value == "KEEP PRC EXPIRY"
        assert workbook["PRC"]["A12"].value == expected["PRC"][1]

        for sheet_name in ("CIV", "MEC", "ELE", "PRC"):
            formulas = [
                str(validation.formula1 or "")
                for validation in workbook[sheet_name].data_validations.dataValidation
            ]
            assert any("OR(A10=" in formula for formula in formulas)
            assert any("INDIRECT(B10)" in formula for formula in formulas)
    finally:
        workbook.close()

    ele_preview = preview_plan_table(revised, "ELE")
    mec_preview = preview_plan_table(revised, "MEC")
    prc_preview = preview_plan_table(revised, "PRC")
    assert ele_preview["count"] == 2
    assert prc_preview["count"] == 2
    assert ele_preview["fromWorkbook"] is True
    assert prc_preview["fromWorkbook"] is True
    assert mec_preview["count"] == 1
    assert mec_preview["fromWorkbook"] is True
    assert mec_preview["fields"][7] == "rowNo"
    assert mec_preview["headers"][7] == "ROW No."
    assert mec_preview["rows"][0][7] == "MEC-ROW-01"

    undone = fsapply.undo(result.journal_id)
    assert undone.failed == 0
    assert fsapply.file_sha256(path) == original
