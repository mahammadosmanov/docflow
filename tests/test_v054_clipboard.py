import re

from openpyxl import Workbook
from openpyxl.worksheet._read_only import ReadOnlyWorksheet

from docflow.clipboard import (
    EXCEL_CLIPBOARD_SCRIPT,
    _cf_html,
    _normalize_excel_cf_html,
    _payload,
    preview_plan_table,
)
from docflow.contracts import Plan


def test_mail_table_stops_at_subcontractor_and_has_plain_fallback():
    fields = [
        "rfiNo",
        "forms",
        "subcontractor",
    ]
    plan = Plan(
        op_id="rfi.log_add",
        meta={
            "editor": {
                "kind": "recordTable",
                "copyEnabled": True,
                "copyFields": fields,
                "copyHeaders": ["RFI No.", "FORMS", "SUBCONTRACTOR"],
                "rows": [
                    {
                        "rfiNo": "AZE-SUN-INT-QMS-RFI-CIV-0860",
                        "forms": "CON-CIV-001",
                        "subcontractor": "AZFEN",
                        "expiry": "MUST NOT COPY",
                        "status": "MUST NOT COPY",
                    }
                ],
            }
        },
    )

    fragment, plain, count = _payload(plan)
    assert count == 1
    assert "AZE-SUN-INT-QMS-RFI-CIV-0860" in fragment
    assert "\t" in plain
    assert "MUST NOT COPY" not in fragment
    assert "MUST NOT COPY" not in plain

    encoded = _cf_html(fragment)
    assert b"StartFragment:" in encoded
    assert b"<!--StartFragment-->" in encoded
    assert encoded.endswith(b"\0")


def test_cf_html_offsets_are_utf8_byte_offsets():
    encoded = _cf_html("<table><tr><td>Şimal sahəsi</td></tr></table>")
    fields = {
        key.decode(): int(value)
        for key, value in re.findall(
            rb"(StartHTML|EndHTML|StartFragment|EndFragment):(\d+)",
            encoded,
        )
    }
    assert encoded[fields["StartHTML"] :].startswith(b"<html>")
    assert encoded[fields["StartFragment"] :].startswith(b"<table>")
    assert encoded[fields["EndFragment"] :].startswith(b"<!--EndFragment-->")
    assert encoded[fields["EndHTML"] :].startswith(b"\0")


def test_excel_cf_html_normalization_preserves_css_and_repairs_offsets():
    document = (
        "<html><head><style>.xl65{font-weight:700;color:#fff}</style></head>"
        "<body><!--StartFragment--><table><tr>"
        "<td class=\"xl65\">Şimal sahəsi</td>"
        "</tr></table><!--EndFragment--></body></html>"
    )
    raw = (
        b"Version:0.9\r\nStartHTML:0000000000\r\nEndHTML:0000000000\r\n"
        b"StartFragment:0000000000\r\nEndFragment:0000000000\r\n"
        + document.encode()
    )

    encoded = _normalize_excel_cf_html(raw)
    fields = {
        key.decode(): int(value)
        for key, value in re.findall(
            rb"(StartHTML|EndHTML|StartFragment|EndFragment):(\d+)",
            encoded,
        )
    }

    assert b".xl65{font-weight:700;color:#fff}" in encoded
    assert encoded[fields["StartHTML"] :].startswith(b"<html>")
    assert encoded[fields["StartFragment"] :].startswith(b"<table>")
    assert encoded[fields["EndFragment"] :].startswith(b"<!--EndFragment-->")
    assert encoded[fields["EndHTML"] :].startswith(b"\0")


def test_excel_copy_uses_visible_native_excel_selection():
    hidden = EXCEL_CLIPBOARD_SCRIPT.index("$excel.Visible = $false")
    copied = EXCEL_CLIPBOARD_SCRIPT.index("$null = $copyRange.Copy()")
    shown = EXCEL_CLIPBOARD_SCRIPT.index("$excel.Visible = $true")
    ready = EXCEL_CLIPBOARD_SCRIPT.index("Set-DocFlowPhase 'ready'")
    assert hidden < copied < shown < ready
    assert "$excel.Interactive = $false" in EXCEL_CLIPBOARD_SCRIPT
    assert "$excel.Interactive = $true" in EXCEL_CLIPBOARD_SCRIPT
    assert "$criteria -contains $rfi" in EXCEL_CLIPBOARD_SCRIPT
    assert "$matchedRows.Add($sourceRow)" in EXCEL_CLIPBOARD_SCRIPT
    assert "$filterRange.AutoFilter" in EXCEL_CLIPBOARD_SCRIPT
    assert ").Hidden = $true" in EXCEL_CLIPBOARD_SCRIPT
    assert ").Hidden = $false" in EXCEL_CLIPBOARD_SCRIPT
    assert "$filterRange.SpecialCells(12)" in EXCEL_CLIPBOARD_SCRIPT
    assert "$copyRange.Select()" in EXCEL_CLIPBOARD_SCRIPT
    assert "$copyRange.Copy()" in EXCEL_CLIPBOARD_SCRIPT
    assert "DOCFLOW_EXCEL_COPY_OK" in EXCEL_CLIPBOARD_SCRIPT
    assert "DocFlowClipboardReader" not in EXCEL_CLIPBOARD_SCRIPT
    assert "Clipboard]::Set" not in EXCEL_CLIPBOARD_SCRIPT
    assert "GetWindowThreadProcessId" in EXCEL_CLIPBOARD_SCRIPT
    assert "Stop-Process -Id $excelPid" in EXCEL_CLIPBOARD_SCRIPT
    assert "FinalReleaseComObject" in EXCEL_CLIPBOARD_SCRIPT


def test_mail_table_reads_applied_values_and_headers_from_workbook(tmp_path):
    path = tmp_path / "RFI Log.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "CIV"
    sheet["A9"] = "RFI No. FROM EXCEL"
    sheet["E9"] = "FORMS FROM EXCEL"
    sheet["O9"] = "SUBCONTRACTOR FROM EXCEL"
    sheet["A10"] = "AZE-SUN-INT-QMS-RFI-CIV-0860"
    sheet["E10"] = "ACTUAL FORM"
    sheet["O10"] = "ACTUAL COMPANY"
    sheet.column_dimensions["A"].width = 25
    workbook.save(path)
    workbook.close()

    plan = Plan(
        op_id="rfi.log_add",
        meta={
            "editor": {
                "kind": "recordTable",
                "copyEnabled": True,
                "copyFields": ["rfiNo", "forms", "subcontractor"],
                "copyHeaders": ["preview rfi", "preview forms", "preview company"],
                "copySource": {
                    "workbook": str(path),
                    "sheet": "CIV",
                    "headerRow": 9,
                    "startRow": 10,
                    "rfiNumbers": ["AZE-SUN-INT-QMS-RFI-CIV-0860"],
                    "columns": [1, 5, 15],
                },
                "rows": [
                    {
                        "rfiNo": "AZE-SUN-INT-QMS-RFI-CIV-0860",
                        "forms": "PREVIEW FORM",
                        "subcontractor": "PREVIEW COMPANY",
                    }
                ],
            }
        },
    )

    fragment, plain, count = _payload(plan)

    assert count == 1
    assert "RFI No. FROM EXCEL" in fragment
    assert "ACTUAL FORM" in fragment
    assert "ACTUAL COMPANY" in plain
    assert "PREVIEW FORM" not in fragment

    preview = preview_plan_table(plan)
    assert preview["headers"] == [
        "RFI No. FROM EXCEL",
        "FORMS FROM EXCEL",
        "SUBCONTRACTOR FROM EXCEL",
    ]
    assert preview["rows"] == [[
        "AZE-SUN-INT-QMS-RFI-CIV-0860",
        "ACTUAL FORM",
        "ACTUAL COMPANY",
    ]]
    assert preview["fromWorkbook"] is True
    assert preview["count"] == 1
    assert preview["columnWidths"][0] == 25


def test_mail_table_finds_filtered_rfis_after_workbook_rows_are_reordered(tmp_path):
    path = tmp_path / "RFI Log.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "CIV"
    sheet["A9"] = "RFI No."
    sheet["E9"] = "FORMS"
    sheet["O9"] = "SUBCONTRACTOR"
    sheet["A10"] = "AZE-SUN-INT-QMS-RFI-CIV-0100"
    sheet["E10"] = "OLD"
    sheet["A20"] = "AZE-SUN-INT-QMS-RFI-CIV-0861"
    sheet["E20"] = "SECOND"
    sheet["O20"] = "AZFEN"
    sheet["A30"] = "AZE-SUN-INT-QMS-RFI-CIV-0860"
    sheet["E30"] = "FIRST"
    sheet["O30"] = "AZFEN"
    workbook.save(path)
    workbook.close()

    plan = Plan(
        op_id="rfi.log_add",
        meta={
            "editor": {
                "kind": "recordTable",
                "copyEnabled": True,
                "copyFields": ["rfiNo", "forms", "subcontractor"],
                "copyHeaders": ["RFI No.", "FORMS", "SUBCONTRACTOR"],
                "copySource": {
                    "workbook": str(path),
                    "sheet": "CIV",
                    "headerRow": 9,
                    "startRow": 10,
                    "rfiNumbers": [
                        "AZE-SUN-INT-QMS-RFI-CIV-0860",
                        "AZE-SUN-INT-QMS-RFI-CIV-0861",
                    ],
                    "columns": [1, 5, 15],
                },
                "rows": [
                    {"rfiNo": "AZE-SUN-INT-QMS-RFI-CIV-0860"},
                    {"rfiNo": "AZE-SUN-INT-QMS-RFI-CIV-0861"},
                ],
            }
        },
    )

    fragment, plain, count = _payload(plan)

    assert count == 2
    assert "FIRST" in fragment
    assert "SECOND" in fragment
    assert "OLD" not in fragment
    assert plain.index("0861") < plain.index("0860")


def test_mail_table_streams_large_read_only_sheet_without_random_cell_reads(
    tmp_path,
    monkeypatch,
):
    path = tmp_path / "RFI Log.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "CIV"
    sheet["A9"] = "RFI No."
    sheet["O9"] = "SUBCONTRACTOR"
    sheet["A5000"] = "AZE-SUN-INT-QMS-RFI-CIV-0860"
    sheet["O5000"] = "AZFEN"
    workbook.save(path)
    workbook.close()

    original_cell = ReadOnlyWorksheet.cell

    def reject_data_random_access(self, row, column, value=None):
        if row > 9:
            raise AssertionError("read-only data rows must be streamed")
        return original_cell(self, row, column, value)

    monkeypatch.setattr(ReadOnlyWorksheet, "cell", reject_data_random_access)
    plan = Plan(
        op_id="rfi.log_add",
        meta={
            "editor": {
                "kind": "recordTable",
                "copyEnabled": True,
                "copyFields": ["rfiNo", "subcontractor"],
                "copyHeaders": ["RFI No.", "SUBCONTRACTOR"],
                "copySource": {
                    "workbook": str(path),
                    "sheet": "CIV",
                    "headerRow": 9,
                    "startRow": 5000,
                    "rfiNumbers": ["AZE-SUN-INT-QMS-RFI-CIV-0860"],
                    "columns": [1, 15],
                },
                "rows": [{"rfiNo": "AZE-SUN-INT-QMS-RFI-CIV-0860"}],
            }
        },
    )

    fragment, plain, count = _payload(plan)

    assert count == 1
    assert "AZE-SUN-INT-QMS-RFI-CIV-0860" in fragment
    assert "AZFEN" in plain
