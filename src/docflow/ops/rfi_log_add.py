"""Bir veya daha fazla RFI Log sheet'inde kayit olusturma.

Is kurallari preview'da hesaplanir. Apply yalnizca onaylanan kayitlari, onceden
belirlenen satirlara yazar; kaynak dosyanin preview'dan sonra degismedigini
SHA-256 ile dogrular. Expiry/Status/Validity sutunlarina dokunulmaz.
"""

from __future__ import annotations

import copy
import ctypes
import os
import pathlib
import re
import shutil
import time as time_module
import unicodedata
import uuid
import warnings as pywarnings
from dataclasses import dataclass
from datetime import date, datetime, time, timedelta
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from ..contracts import (
    BoolParam,
    ChoiceParam,
    FileParam,
    LongTextParam,
    JournalEntry,
    Operation,
    ParamSchema,
    Plan,
    Result,
    Step,
    SecretParam,
    SheetCountsParam,
    Warning,
)
from ..email_extract import detect_discipline, extract_email
from ..fs.apply import file_sha256, write_journal
from ..registry import register
from ..progress import cancellation_requested
from ..rpc import notify
from ..xlsx_ooxml import patch_workbook_sheets

pywarnings.filterwarnings(
    "ignore",
    message="Data Validation extension is not supported.*",
    category=UserWarning,
    module=r"openpyxl\..*",
)

SUPPORTED_DISCIPLINES = ("CIV", "MEC", "ELE", "PRC")
RFI_RE = re.compile(r"-RFI-(CIV|MEC|ELE|PRC)-(\d+)$", re.IGNORECASE)
ITEM_RE = re.compile(r"^(CIV|MEC|ELE|PRC)-(\d+)\.", re.IGNORECASE)
NORTH_RE = re.compile(r"(?:^|[^A-Z0-9])N\d{1,3}(?:[^A-Z0-9]|$)", re.IGNORECASE)
SOUTH_RE = re.compile(r"(?:^|[^A-Z0-9])S\d{1,3}(?:[^A-Z0-9]|$)", re.IGNORECASE)
DAY_START = time(9, 0)
DAY_END = time(17, 30)
SLOT = timedelta(minutes=30)
MAX_RFI_COUNT = 100


class WorkbookLockedError(RuntimeError):
    """Raised when Windows is keeping the RFI workbook open or read-only."""


def _default_workbook() -> str | None:
    configured = os.environ.get("DOCFLOW_RFI_LOG")
    if configured:
        return configured
    desktop = pathlib.Path.home() / "OneDrive - AZFEN LLC" / "Desktop"
    candidate = desktop / "SystemU" / "RFI Log.xlsx"
    if candidate.is_file():
        return str(candidate)

    # OneDrive may relocate Desktop folders beneath a user-created container
    # (for example Desktop/Rohtolos/SystemU) while preserving the workbook.
    relocated = sorted(desktop.glob("*/SystemU/RFI Log.xlsx"))
    return str(relocated[0]) if relocated else None


@dataclass(frozen=True)
class Layout:
    sheet: str
    issue_col: int
    inspection_col: int
    expiry_col: int
    time_col: int
    azfen_col: int
    intec_col: int
    location_col: int
    subcontractor_col: int
    status_col: int
    validity_col: int
    row_col: int | None = None


LAYOUTS = {
    "CIV": Layout("CIV", 8, 9, 10, 11, 12, 13, 14, 15, 16, 17),
    "MEC": Layout("MEC", 9, 10, 11, 12, 13, 14, 15, 16, 17, 18, row_col=8),
    "ELE": Layout("ELE", 8, 9, 10, 11, 12, 13, 14, 15, 16, 17),
    "PRC": Layout("PRC", 8, 9, 10, 11, 12, 13, 14, 15, 16, 17),
}

COPY_FIELDS = [
    "rfiNo",
    "itpNo",
    "itemNo",
    "description",
    "forms",
    "tagNo",
    "area",
    "dateOfIssue",
    "inspectionDate",
    "startTime",
    "azfenType",
    "intecType",
    "location",
    "subcontractor",
]

COPY_HEADERS = [
    "RFI No.",
    "ITP No.",
    "Item No.",
    "INSPECTION DESCRIPTION",
    "FORMS",
    "TAG NO",
    "AREA",
    "DATE OF ISSUE",
    "INSPECTION DATE",
    "INSPECTION START TIME",
    "INSP. TYPE AZFEN",
    "INSP. TYPE INTEC",
    "Inspection Location",
    "SUBCONTRACTOR",
]

EDITABLE_FIELDS = [*COPY_FIELDS, "rowNo"]


def _today() -> date:
    return date.today()


def _signature(path: pathlib.Path) -> dict[str, Any]:
    stat = path.stat()
    return {
        "size": stat.st_size,
        "sha256": file_sha256(path),
    }


def _as_date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return None


def _as_time(value: Any) -> time | None:
    if isinstance(value, datetime):
        return value.time().replace(second=0, microsecond=0)
    if isinstance(value, time):
        return value.replace(second=0, microsecond=0)
    if isinstance(value, (int, float)):
        seconds = round(float(value) * 24 * 60 * 60)
        return time((seconds // 3600) % 24, (seconds % 3600) // 60)
    if isinstance(value, str):
        try:
            return datetime.strptime(value.strip(), "%H:%M").time()
        except ValueError:
            return None
    return None


def _date_text(value: date) -> str:
    return value.strftime("%d-%b-%y")


def _time_text(value: time) -> str:
    return value.strftime("%H:%M")


def _parse_date_text(value: str) -> date:
    for fmt in ("%d-%b-%y", "%Y-%m-%d", "%d.%m.%Y"):
        try:
            return datetime.strptime(value.strip(), fmt).date()
        except ValueError:
            continue
    raise ValueError(f"gecersiz tarih: {value}")


def _parse_time_text(value: str) -> time:
    try:
        return datetime.strptime(value.strip(), "%H:%M").time()
    except ValueError as exc:
        raise ValueError(f"gecersiz saat: {value}") from exc


def _canonical_itp(discipline: str, number: int) -> str:
    return f"AZE-SUN-INT-QMS-ITP-{discipline}-{number:04d}"


def _itp_signature(value: str) -> tuple[str, ...]:
    """Compare ITP identifiers by tokens, not their often-corrupted separators."""

    normalized = unicodedata.normalize("NFKC", value).upper()
    return tuple(re.findall(r"[A-Z]+|\d+", normalized))


def _resolve_itp(
    value: str,
    item_no: str,
    itp_to_items: dict[str, list[str]],
) -> str | None:
    if value in itp_to_items:
        return value

    signature = _itp_signature(value)
    signature_matches = [
        label
        for label in itp_to_items
        if signature and _itp_signature(label) == signature
    ]
    if len(signature_matches) == 1:
        return signature_matches[0]

    # The selected Item No. already has a definitive ITP relationship. This
    # also repairs old saved rows whose separator was decoded as "Ë—".
    item_matches = [
        label
        for label, items in itp_to_items.items()
        if item_no in items
    ]
    return item_matches[0] if len(item_matches) == 1 else None


def _recent_unique(values: list[str], limit: int = 500) -> list[str]:
    seen: set[str] = set()
    out: list[str] = []
    for value in reversed(values):
        clean = value.strip()
        if not clean or clean in seen:
            continue
        seen.add(clean)
        out.append(clean)
        if len(out) >= limit:
            break
    return out


def _infer_area(tag: str) -> str:
    leading = re.match(r"^\s*([NS])\d{1,3}(?:[^A-Z0-9]|$)", tag, re.IGNORECASE)
    if leading:
        return "North Site" if leading.group(1).upper() == "N" else "South Site"
    north = bool(NORTH_RE.search(tag))
    south = bool(SOUTH_RE.search(tag))
    if north == south:
        return ""
    return "North Site" if north else "South Site"


def _header_key(value: Any) -> str:
    normalized = unicodedata.normalize("NFKD", str(value or "")).casefold()
    return "".join(char for char in normalized if char.isalnum())


def _mojibake_score(value: str) -> int:
    return sum(value.count(marker) for marker in ("Ã", "Â", "â", "ð", "Ë"))


def _clean_workbook_text(value: Any) -> str:
    """Normalize workbook/user text and repair UTF-8 decoded as CP1252.

    Repair is deliberately conservative: the candidate must be valid UTF-8
    and contain fewer well-known mojibake lead characters. Correct Unicode
    text, including U+2022 bullets, is returned unchanged.
    """

    text = unicodedata.normalize("NFC", str(value or "").strip())
    for _attempt in range(2):
        if _mojibake_score(text) == 0:
            break
        try:
            repaired = text.encode("cp1252").decode("utf-8")
        except (UnicodeEncodeError, UnicodeDecodeError):
            break
        if _mojibake_score(repaired) >= _mojibake_score(text):
            break
        text = unicodedata.normalize("NFC", repaired)
    return text


def _lookup_columns(lookup_sheet: Any) -> dict[str, int]:
    """Locate the seven-column Item/description/forms/type lookup block.

    The production workbook has gained an extra ITP-list column over time,
    shifting this block from K:Q to L:R.  Header-based discovery keeps both
    workbook layouts valid and prevents formulas from pointing at stale
    columns.
    """

    headers = list(
        next(
            lookup_sheet.iter_rows(
                min_row=1,
                max_row=1,
                min_col=1,
                max_col=lookup_sheet.max_column,
                values_only=True,
            )
        )
    )
    item_headers = {"item", "itemno", "itemnumber"}
    for offset, value in enumerate(headers):
        if _header_key(value) not in item_headers or offset + 6 >= len(headers):
            continue
        if (
            _header_key(headers[offset + 3]) not in item_headers
            or _header_key(headers[offset + 5]) not in item_headers
        ):
            continue
        start = offset + 1
        return {
            "item": start,
            "description": start + 1,
            "forms": start + 2,
            "azfenItem": start + 3,
            "azfenType": start + 4,
            "intecItem": start + 5,
            "intecType": start + 6,
        }
    raise ValueError("Sheet1 Item/description/forms/type esleme kolonlari bulunamadi")


def _read_source(
    path: pathlib.Path,
    discipline: str,
    *,
    workbook: Any | None = None,
) -> dict[str, Any]:
    layout = LAYOUTS[discipline]
    owns_workbook = workbook is None
    if workbook is None:
        with pywarnings.catch_warnings():
            pywarnings.filterwarnings(
                "ignore",
                message="Data Validation extension is not supported.*",
                category=UserWarning,
            )
            workbook = load_workbook(
                path,
                read_only=True,
                data_only=True,
                keep_links=True,
            )
    try:
        if layout.sheet not in workbook.sheetnames or "Sheet1" not in workbook.sheetnames:
            raise ValueError(
                f"RFI Log icin {layout.sheet}/Sheet1 sayfalari bulunamadi"
            )
        sheet = workbook[layout.sheet]
        lookup_sheet = workbook["Sheet1"]
        lookup_columns = _lookup_columns(lookup_sheet)

        presets: dict[str, dict[str, str]] = {}
        for row in lookup_sheet.iter_rows(
            min_row=1,
            max_row=lookup_sheet.max_row,
            min_col=lookup_columns["item"],
            max_col=lookup_columns["intecType"],
            values_only=True,
        ):
            item = _clean_workbook_text(row[0])
            match = ITEM_RE.match(item)
            if not match or match.group(1).upper() != discipline:
                continue
            azfen_item = _clean_workbook_text(row[3])
            intec_item = _clean_workbook_text(row[5])
            if azfen_item != item or intec_item != item:
                raise ValueError(
                    f"{discipline} Sheet1 Item eslemesi tutarsiz: {item}"
                )
            preset = {
                "description": _clean_workbook_text(row[1]),
                "forms": _clean_workbook_text(row[2]),
                "azfenType": _clean_workbook_text(row[4]),
                "intecType": _clean_workbook_text(row[6]),
            }
            missing = [field for field, value in preset.items() if not value]
            if missing:
                raise ValueError(
                    f"{discipline} Sheet1 esleme verisi eksik: "
                    f"{item} / {', '.join(missing)}"
                )
            existing = presets.get(item)
            if existing is not None and existing != preset:
                raise ValueError(
                    f"{discipline} Sheet1 cakisan Item eslemesi: {item}"
                )
            presets[item] = preset

        max_number = 0
        last_row = 9
        latest_item = ""
        tags: list[str] = []
        areas: list[str] = []
        locations: list[str] = []
        subcontractors: list[str] = []
        occupied: dict[date, dict[str, time]] = {}

        for row_number, row in enumerate(
            sheet.iter_rows(
            min_row=10,
            max_row=sheet.max_row,
            min_col=1,
            max_col=layout.validity_col,
            values_only=True,
            ),
            start=10,
        ):
            rfi = str(row[0] or "").strip()
            match = RFI_RE.search(rfi)
            if not match or match.group(1).upper() != discipline:
                continue
            last_row = row_number
            max_number = max(max_number, int(match.group(2)))
            latest_item = _clean_workbook_text(row[2]) or latest_item
            tags.append(_clean_workbook_text(row[5]))
            areas.append(_clean_workbook_text(row[6]))
            locations.append(_clean_workbook_text(row[layout.location_col - 1]))
            subcontractors.append(
                _clean_workbook_text(row[layout.subcontractor_col - 1])
            )
            inspection_day = _as_date(row[layout.inspection_col - 1])
            inspection_time = _as_time(row[layout.time_col - 1])
            if inspection_day and inspection_time:
                occupied.setdefault(inspection_day, {})[rfi] = inspection_time

        if not presets:
            raise ValueError(f"{discipline} icin Sheet1 esleme verisi bulunamadi")

        workbook_itps: dict[int, str] = {}
        for value in next(
            lookup_sheet.iter_rows(min_row=1, max_row=1, values_only=True)
        ):
            label = str(value or "").strip()
            if "ITP" not in label.upper() or discipline not in label.upper():
                continue
            number_match = re.search(r"(\d{4})$", label)
            if number_match:
                workbook_itps[int(number_match.group(1))] = label

        itp_to_items: dict[str, list[str]] = {}
        for item in presets:
            match = ITEM_RE.match(item)
            assert match is not None
            number = int(match.group(2))
            itp = workbook_itps.get(number, _canonical_itp(discipline, number))
            itp_to_items.setdefault(itp, []).append(item)
        for items in itp_to_items.values():
            items.sort(
                key=lambda value: [
                    int(part) if part.isdigit() else part
                    for part in re.split(r"(\d+)", value)
                ]
            )

        default_item = latest_item if latest_item in presets else next(iter(presets))
        item_match = ITEM_RE.match(default_item)
        assert item_match is not None
        default_number = int(item_match.group(2))
        default_itp = workbook_itps.get(
            default_number,
            _canonical_itp(discipline, default_number),
        )
        return {
            "presets": presets,
            "itpToItems": itp_to_items,
            "lookupColumns": lookup_columns,
            "defaultItem": default_item,
            "defaultItp": default_itp,
            "maxNumber": max_number,
            "lastRow": last_row,
            "maxRow": sheet.max_row,
            "tags": _recent_unique(tags),
            "areas": _recent_unique(areas),
            "locations": _recent_unique(locations),
            "subcontractors": _recent_unique(subcontractors),
            "occupied": occupied,
        }
    finally:
        if owns_workbook:
            workbook.close()


def _read_sources(
    path: pathlib.Path,
    requested: list[str],
) -> dict[str, dict[str, Any]]:
    """Read requested sheet data once; each sheet keeps its own clock."""

    with pywarnings.catch_warnings():
        pywarnings.filterwarnings(
            "ignore",
            message="Data Validation extension is not supported.*",
            category=UserWarning,
        )
        workbook = load_workbook(
            path,
            read_only=True,
            data_only=True,
            keep_links=True,
        )
    try:
        missing = [
            discipline
            for discipline in requested
            if LAYOUTS[discipline].sheet not in workbook.sheetnames
        ]
        if missing:
            raise ValueError(
                "RFI Log'da secilen sheet bulunamadi: " + ", ".join(missing)
            )

        return {
            discipline: _read_source(path, discipline, workbook=workbook)
            for discipline in requested
        }
    finally:
        workbook.close()


def _next_slot(
    occupied: dict[date, dict[str, time]],
    day: date,
) -> time:
    times = list(occupied.get(day, {}).values())
    if not times:
        return DAY_START
    last = max(times)
    return (datetime.combine(day, last) + SLOT).time()


def _allocate(
    occupied: dict[date, dict[str, time]],
    start_day: date,
    count: int,
    wrap: bool,
) -> tuple[list[tuple[date, time]], int]:
    working = {day: dict(values) for day, values in occupied.items()}
    slots: list[tuple[date, time]] = []
    overflowed = 0
    day = start_day
    for index in range(count):
        candidate = _next_slot(working, day)
        if candidate > DAY_END:
            overflowed += 1
            if wrap:
                while candidate > DAY_END:
                    day += timedelta(days=1)
                    candidate = _next_slot(working, day)
        marker = f"__planned_{index}"
        working.setdefault(day, {})[marker] = candidate
        slots.append((day, candidate))
    return slots, overflowed


def _columns() -> list[dict[str, Any]]:
    return [
        {
            "key": "_sheet",
            "labelKey": "rfi.col.sheet",
            "kind": "readonly",
            "width": 48,
            "compact": True,
        },
        {
            "key": "rfiNo",
            "labelKey": "rfi.col.rfiNo",
            "kind": "readonly",
            "width": 72,
            "compact": True,
            "display": "suffix",
        },
        {
            "key": "itpNo",
            "labelKey": "rfi.col.itpNo",
            "kind": "choice",
            "width": 72,
            "compact": True,
            "display": "suffix",
        },
        {
            "key": "itemNo",
            "labelKey": "rfi.col.itemNo",
            "kind": "choice",
            "width": 76,
            "compact": True,
            "display": "itemSuffix",
        },
        {
            "key": "description",
            "labelKey": "rfi.col.description",
            "kind": "multiline",
            "width": 360,
            "detail": True,
        },
        {
            "key": "forms",
            "labelKey": "rfi.col.forms",
            "kind": "multiline",
            "width": 280,
            "detail": True,
        },
        {
            "key": "tagNo",
            "labelKey": "rfi.col.tagNo",
            "kind": "autocomplete",
            "width": 116,
            "compact": True,
        },
        {
            "key": "area",
            "labelKey": "rfi.col.area",
            "kind": "autocomplete",
            "width": 116,
            "compact": True,
        },
        {
            "key": "rowNo",
            "labelKey": "rfi.col.rowNo",
            "kind": "text",
            "width": 82,
            "compact": True,
            "sheets": ["MEC"],
        },
        {
            "key": "dateOfIssue",
            "labelKey": "rfi.col.issueDate",
            "kind": "text",
            "width": 90,
            "compact": True,
        },
        {
            "key": "inspectionDate",
            "labelKey": "rfi.col.inspectionDate",
            "kind": "text",
            "width": 94,
            "compact": True,
        },
        {
            "key": "startTime",
            "labelKey": "rfi.col.startTime",
            "kind": "text",
            "width": 68,
            "compact": True,
        },
        {
            "key": "azfenType",
            "labelKey": "rfi.col.azfenType",
            "kind": "text",
            "width": 130,
            "detail": True,
        },
        {
            "key": "intecType",
            "labelKey": "rfi.col.intecType",
            "kind": "text",
            "width": 130,
            "detail": True,
        },
        {
            "key": "location",
            "labelKey": "rfi.col.location",
            "kind": "autocomplete",
            "width": 134,
            "compact": True,
        },
        {
            "key": "subcontractor",
            "labelKey": "rfi.col.subcontractor",
            "kind": "autocomplete",
            "width": 110,
            "compact": True,
        },
    ]


def _plan_meta(
    sources: dict[str, dict[str, Any]],
    rows: list[dict[str, Any]],
) -> dict[str, Any]:
    itp_to_items: dict[str, list[str]] = {}
    presets: dict[str, dict[str, str]] = {}
    autocomplete: dict[str, list[str]] = {
        "tagNo": [],
        "area": [],
        "location": [],
        "subcontractor": ["AZFEN", "PDL"],
    }
    for discipline in SUPPORTED_DISCIPLINES:
        source = sources.get(discipline)
        if source is None:
            continue
        itp_to_items.update(source["itpToItems"])
        presets.update(source["presets"])
        for field, source_key in (
            ("tagNo", "tags"),
            ("area", "areas"),
            ("location", "locations"),
            ("subcontractor", "subcontractors"),
        ):
            autocomplete[field] = list(
                dict.fromkeys(
                    [*autocomplete[field], *list(source[source_key])]
                )
            )
    return {
        "editor": {
            "kind": "recordTable",
            "layout": "compact",
            "titleKey": "rfi.editor.title",
            "groupField": "rfiNo",
            "selectionGroupFields": ["_sheet", "rfiNo"],
            "columns": _columns(),
            "rows": copy.deepcopy(rows),
            "choices": {
                "itpNo": sorted(itp_to_items),
            },
            "choicesByValue": {
                "itemNo": {
                    "dependsOn": "itpNo",
                    "values": itp_to_items,
                },
            },
            "autoFill": {
                "field": "itemNo",
                "values": presets,
            },
            "autocomplete": autocomplete,
            "regexFill": [
                {
                    "source": "tagNo",
                    "target": "area",
                    "pattern": NORTH_RE.pattern,
                    "value": "North Site",
                },
                {
                    "source": "tagNo",
                    "target": "area",
                    "pattern": SOUTH_RE.pattern,
                    "value": "South Site",
                },
            ],
            "duplicateEnabled": True,
            "selectionEnabled": True,
            "bulkEnabled": True,
            "detailsEnabled": True,
            "detailsDefaultVisible": False,
            "copyEnabled": True,
            "copyFields": COPY_FIELDS,
            "copyHeaders": COPY_HEADERS,
            "sheets": [
                discipline
                for discipline in SUPPORTED_DISCIPLINES
                if discipline in sources
            ],
            "sheetCounts": {
                discipline: sum(
                    1 for row in rows if row.get("_sheet") == discipline
                )
                for discipline in sources
            },
        }
    }


def _copy_source(
    workbook_path: pathlib.Path,
    discipline: str,
    start_row: int,
    rows: list[dict[str, Any]],
) -> dict[str, Any]:
    layout = LAYOUTS[discipline]
    fields = list(COPY_FIELDS)
    headers = list(COPY_HEADERS)
    columns = [1, 2, 3, 4, 5, 6, 7]
    if layout.row_col is not None:
        fields.insert(7, "rowNo")
        headers.insert(7, "ROW No.")
        columns.append(layout.row_col)
    columns.extend(
        [
            layout.issue_col,
            layout.inspection_col,
            layout.time_col,
            layout.azfen_col,
            layout.intec_col,
            layout.location_col,
            layout.subcontractor_col,
        ]
    )
    return {
        "workbook": str(workbook_path),
        "sheet": layout.sheet,
        "headerRow": 9,
        "startRow": start_row,
        "rfiNumbers": [str(row.get("rfiNo") or "") for row in rows],
        "fields": fields,
        "headers": headers,
        "columns": columns,
    }


def _record(
    *,
    index: int,
    discipline: str,
    rfi: str,
    source: dict[str, Any],
    issue_day: date,
    slot: tuple[date, time],
) -> dict[str, Any]:
    preset = source["presets"][source["defaultItem"]]
    return {
        "_rowId": f"{discipline}-g{index + 1}-r1",
        "_sheet": discipline,
        "rfiNo": rfi,
        "itpNo": source["defaultItp"],
        "itemNo": source["defaultItem"],
        "description": preset["description"],
        "forms": preset["forms"],
        "tagNo": "",
        "area": "",
        "rowNo": "",
        "dateOfIssue": _date_text(issue_day),
        "inspectionDate": _date_text(slot[0]),
        "startTime": _time_text(slot[1]),
        "azfenType": preset["azfenType"],
        "intecType": preset["intecType"],
        "location": "",
        "subcontractor": "AZFEN",
    }


def _apply_email_values(
    row: dict[str, Any],
    extracted: dict[str, Any],
    source: dict[str, Any],
    notes: list[str],
) -> None:
    item = str(extracted.get("itemNo") or "").strip()
    itp = str(extracted.get("itpNo") or "").strip()
    if item:
        if item in source["presets"]:
            preset = source["presets"][item]
            row["itemNo"] = item
            resolved = _resolve_itp(itp, item, source["itpToItems"])
            if resolved is None:
                matches = [
                    label
                    for label, items in source["itpToItems"].items()
                    if item in items
                ]
                resolved = matches[0] if len(matches) == 1 else None
            if resolved is not None:
                row["itpNo"] = resolved
            row.update(
                {
                    "description": preset["description"],
                    "forms": preset["forms"],
                    "azfenType": preset["azfenType"],
                    "intecType": preset["intecType"],
                }
            )
        else:
            notes.append(f"Unknown Item No. ignored: {item}")
    elif itp:
        resolved = _resolve_itp(itp, row["itemNo"], source["itpToItems"])
        if resolved is not None:
            row["itpNo"] = resolved
            allowed = source["itpToItems"][resolved]
            if row["itemNo"] not in allowed and len(allowed) == 1:
                item = allowed[0]
                preset = source["presets"][item]
                row.update(
                    {
                        "itemNo": item,
                        "description": preset["description"],
                        "forms": preset["forms"],
                        "azfenType": preset["azfenType"],
                        "intecType": preset["intecType"],
                    }
                )
        else:
            notes.append(f"Unknown ITP No. ignored: {itp}")

    for field in ("tagNo", "area", "location", "subcontractor"):
        value = str(extracted.get(field) or "").strip()
        if value:
            row[field] = value
    if not row["area"]:
        row["area"] = _infer_area(row["tagNo"])


def _rows_from_email(
    base_rows: list[dict[str, Any]],
    extraction: dict[str, Any],
    source: dict[str, Any],
) -> tuple[list[dict[str, Any]], list[str]]:
    extracted_rows = extraction.get("rows")
    if not isinstance(extracted_rows, list) or not extracted_rows:
        return base_rows, []

    grouped: dict[int, list[dict[str, Any]]] = {}
    for raw in extracted_rows[:MAX_RFI_COUNT]:
        if not isinstance(raw, dict):
            continue
        try:
            index = int(raw.get("rfiIndex") or 1)
        except (TypeError, ValueError):
            continue
        if 1 <= index <= len(base_rows):
            grouped.setdefault(index - 1, []).append(raw)

    rows: list[dict[str, Any]] = []
    notes: list[str] = []
    for index, base in enumerate(base_rows):
        candidates = grouped.get(index) or [{}]
        for duplicate_index, extracted in enumerate(candidates):
            row = copy.deepcopy(base)
            if duplicate_index:
                row["_rowId"] = (
                    f"{row.get('_sheet', '')}-g{index + 1}-"
                    f"r{duplicate_index + 1}"
                )
            _apply_email_values(row, extracted, source, notes)
            rows.append(row)
    return rows, notes


def _warnings_for_rows(
    rows: list[dict[str, Any]],
    *,
    overflowed: int = 0,
    wrap: bool = True,
) -> list[Warning]:
    warnings: list[Warning] = []
    if overflowed:
        key = "warn.rfi.overflowMoved" if wrap else "warn.rfi.overflowKept"
        warnings.append(Warning(key, str(overflowed)))

    by_rfi: dict[str, list[dict[str, Any]]] = {}
    for row in rows:
        by_rfi.setdefault(str(row["rfiNo"]), []).append(row)
    for rfi, group in by_rfi.items():
        schedule = {
            (str(row.get("inspectionDate") or ""), str(row.get("startTime") or ""))
            for row in group
        }
        if len(schedule) > 1:
            warnings.append(Warning("warn.rfi.groupScheduleMismatch", rfi))
    return warnings


def _safe_rows(
    raw_rows: list[Any],
    *,
    discipline: str,
    base_rfis: list[str],
    source: dict[str, Any],
) -> list[dict[str, Any]]:
    layout = LAYOUTS[discipline]
    allowed_fields = set(EDITABLE_FIELDS) | {"_rowId", "_sheet"}
    allowed_rfis = set(base_rfis)
    rows: list[dict[str, Any]] = []
    for index, raw in enumerate(raw_rows):
        if not isinstance(raw, dict):
            raise ValueError("gecersiz RFI satiri")
        row = {key: raw.get(key, "") for key in allowed_fields}
        row["_rowId"] = str(row.get("_rowId") or f"row-{uuid.uuid4().hex[:10]}")
        if str(row.get("_sheet") or "").upper() != discipline:
            raise ValueError(f"RFI satiri yanlis sheet'e ait: {discipline}")
        row["_sheet"] = discipline
        row["rfiNo"] = str(row["rfiNo"]).strip()
        if row["rfiNo"] not in allowed_rfis:
            raise ValueError(f"plan disi RFI numarasi: {row['rfiNo']}")
        for key in EDITABLE_FIELDS:
            row[key] = _clean_workbook_text(row.get(key))
        if layout.row_col is None:
            row["rowNo"] = ""
        if not row["area"]:
            row["area"] = _infer_area(row["tagNo"])
        _parse_date_text(row["dateOfIssue"])
        _parse_date_text(row["inspectionDate"])
        _parse_time_text(row["startTime"])
        if row["itemNo"] not in source["presets"]:
            raise ValueError(f"bilinmeyen Item No.: {row['itemNo']}")
        resolved_itp = _resolve_itp(
            row["itpNo"],
            row["itemNo"],
            source["itpToItems"],
        )
        if resolved_itp is None:
            raise ValueError(f"bilinmeyen ITP No.: {row['itpNo']}")
        row["itpNo"] = resolved_itp
        if row["itemNo"] not in source["itpToItems"][row["itpNo"]]:
            raise ValueError(
                f"Item No. secilen ITP'ye ait degil: {row['itpNo']} / {row['itemNo']}"
            )
        for key in (
            "description",
            "forms",
            "azfenType",
            "intecType",
            "subcontractor",
        ):
            if not row[key]:
                raise ValueError(f"zorunlu alan bos: {key} (satir {index + 1})")
        rows.append(row)

    if not rows:
        raise ValueError("en az bir RFI satiri gerekli")
    present = {row["rfiNo"] for row in rows}
    if present != allowed_rfis:
        missing = sorted(allowed_rfis - present)
        raise ValueError(f"RFI grubu silinemez: {', '.join(missing)}")
    return rows


def _assert_not_locked(path: pathlib.Path) -> None:
    if os.name != "nt":
        return

    from ctypes import wintypes

    kernel32 = ctypes.WinDLL("kernel32", use_last_error=True)
    create_file = kernel32.CreateFileW
    create_file.argtypes = [
        wintypes.LPCWSTR,
        wintypes.DWORD,
        wintypes.DWORD,
        wintypes.LPVOID,
        wintypes.DWORD,
        wintypes.DWORD,
        wintypes.HANDLE,
    ]
    create_file.restype = wintypes.HANDLE
    close_handle = kernel32.CloseHandle
    close_handle.argtypes = [wintypes.HANDLE]
    close_handle.restype = wintypes.BOOL
    invalid_handle = ctypes.c_void_p(-1).value
    last_error: OSError | None = None

    for attempt in range(8):
        handle = create_file(
            str(path),
            0x80000000 | 0x40000000,  # GENERIC_READ | GENERIC_WRITE
            0,  # Excel aciksa paylasimsiz handle basarisiz olur.
            None,
            3,  # OPEN_EXISTING
            0x80,  # FILE_ATTRIBUTE_NORMAL
            None,
        )
        if handle != invalid_handle:
            close_handle(handle)
            return
        last_error = ctypes.WinError(ctypes.get_last_error())
        if attempt < 7:
            time_module.sleep(0.15)
    raise WorkbookLockedError(path.name) from last_error


def _is_workbook_access_error(exc: Exception) -> bool:
    return isinstance(exc, (WorkbookLockedError, PermissionError)) or (
        isinstance(exc, OSError)
        and getattr(exc, "winerror", None) in {5, 32, 33}
    )


def _replace_with_retry(source: pathlib.Path, destination: pathlib.Path) -> None:
    for attempt in range(20):
        try:
            os.replace(source, destination)
            return
        except OSError as exc:
            if not _is_workbook_access_error(exc) or attempt >= 19:
                raise
            time_module.sleep(0.25)


def _copy_contents(
    source: pathlib.Path,
    destination: pathlib.Path,
    *,
    overwrite: bool = False,
) -> None:
    # shutil.copy/copy2/copyfile Windows'ta CopyFile2 kullanir. OneDrive bazen
    # tum byte'lari kopyaladiktan sonra metadata asamasinda EACCES dondurur.
    # Stream kopyasi yalniz workbook icerigini tasir ve bu yalanci hatayi onler.
    mode = "wb" if overwrite else "xb"
    with source.open("rb") as incoming, destination.open(mode) as outgoing:
        shutil.copyfileobj(incoming, outgoing, length=1024 * 1024)


def _apply_error_warning(exc: Exception, workbook_path: pathlib.Path) -> Warning:
    if _is_workbook_access_error(exc):
        return Warning("error.rfi.workbookLocked", workbook_path.name)
    return Warning("error.step_failed", str(exc))


def _formula_or_value(
    value: str,
    default_value: str,
    formula: str,
) -> tuple[str, Any]:
    normalized_value = unicodedata.normalize(
        "NFC",
        value.replace("\r\n", "\n").replace("\r", "\n"),
    )
    normalized_default = unicodedata.normalize(
        "NFC",
        default_value.replace("\r\n", "\n").replace("\r", "\n"),
    )
    return (
        ("formula", (formula, default_value))
        if normalized_value == normalized_default
        else ("value", value)
    )


def _record_updates(
    discipline: str,
    start_row: int,
    rows: list[dict[str, Any]],
    presets: dict[str, dict[str, str]],
    lookup_columns: dict[str, int],
) -> dict[tuple[int, int], tuple[str, Any]]:
    layout = LAYOUTS[discipline]
    updates: dict[tuple[int, int], tuple[str, Any]] = {}
    for index, record in enumerate(rows):
        row_number = start_row + index
        preset = presets[record["itemNo"]]
        values = {
            1: record["rfiNo"],
            2: record["itpNo"],
            3: record["itemNo"],
            6: record["tagNo"],
            7: record["area"],
            layout.issue_col: _parse_date_text(record["dateOfIssue"]),
            layout.inspection_col: _parse_date_text(record["inspectionDate"]),
            layout.time_col: _parse_time_text(record["startTime"]),
            layout.location_col: record["location"],
            layout.subcontractor_col: record["subcontractor"],
        }
        if layout.row_col is not None:
            values[layout.row_col] = record["rowNo"]
        for column, value in values.items():
            updates[(row_number, column)] = ("value", value)

        description_col, forms_col = 4, 5
        azfen_col, intec_col = layout.azfen_col, layout.intec_col

        updates[(row_number, description_col)] = _formula_or_value(
            record["description"],
            preset["description"],
            (
                f'=IFERROR(VLOOKUP(C{row_number},Sheet1!'
                f'{get_column_letter(lookup_columns["item"])}:'
                f'{get_column_letter(lookup_columns["description"])},2,0),"")'
            ),
        )
        updates[(row_number, forms_col)] = _formula_or_value(
            record["forms"],
            preset["forms"],
            (
                f'=IFERROR(VLOOKUP(D{row_number},Sheet1!'
                f'{get_column_letter(lookup_columns["description"])}:'
                f'{get_column_letter(lookup_columns["forms"])},2,0),"")'
            ),
        )
        updates[(row_number, azfen_col)] = _formula_or_value(
            record["azfenType"],
            preset["azfenType"],
            (
                f'=IFERROR(VLOOKUP(C{row_number},Sheet1!'
                f'{get_column_letter(lookup_columns["azfenItem"])}:'
                f'{get_column_letter(lookup_columns["azfenType"])},2,0),"")'
            ),
        )
        updates[(row_number, intec_col)] = _formula_or_value(
            record["intecType"],
            preset["intecType"],
            (
                f'=IFERROR(VLOOKUP(C{row_number},Sheet1!'
                f'{get_column_letter(lookup_columns["intecItem"])}:'
                f'{get_column_letter(lookup_columns["intecType"])},2,0),"")'
            ),
        )
    return updates


def _write_batches(
    workbook_path: pathlib.Path,
    temp_path: pathlib.Path,
    batches: list[dict[str, Any]],
) -> None:
    sheet_updates = {
        str(batch["discipline"]): _record_updates(
            str(batch["discipline"]),
            int(batch["startRow"]),
            list(batch["records"]),
            dict(batch["presets"]),
            dict(batch["lookupColumns"]),
        )
        for batch in batches
    }
    patch_workbook_sheets(
        workbook_path,
        temp_path,
        sheet_updates,
        allow_consecutive_duplicates_column=1,
    )


def _verify_records(
    workbook_path: pathlib.Path,
    discipline: str,
    start_row: int,
    rows: list[dict[str, Any]],
) -> None:
    layout = LAYOUTS[discipline]
    with pywarnings.catch_warnings():
        pywarnings.filterwarnings(
            "ignore",
            message="Data Validation extension is not supported.*",
            category=UserWarning,
        )
        workbook = load_workbook(
            workbook_path,
            read_only=True,
            data_only=False,
            keep_links=True,
        )
    try:
        sheet = workbook[layout.sheet]
        for index, record in enumerate(rows):
            row_number = start_row + index
            checks = {
                1: record["rfiNo"],
                2: record["itpNo"],
                3: record["itemNo"],
                6: record["tagNo"],
                7: record["area"],
                layout.location_col: record["location"],
                layout.subcontractor_col: record["subcontractor"],
            }
            if layout.row_col is not None:
                checks[layout.row_col] = record["rowNo"]
            for column, expected in checks.items():
                actual = str(sheet.cell(row_number, column).value or "")
                if actual != str(expected or ""):
                    raise RuntimeError(
                        f"yazma dogrulamasi basarisiz: {layout.sheet}!"
                        f"{sheet.cell(row_number, column).coordinate}"
                    )
            # Dokunulmamasi gereken kolonlar acikca korunur.
            if layout.expiry_col in (
                layout.issue_col,
                layout.inspection_col,
                layout.time_col,
            ):
                raise AssertionError("gecersiz kolon eslemesi")
    finally:
        workbook.close()


def _requested_sheet_counts(params: dict[str, Any]) -> dict[str, int]:
    raw = params.get("sheet_counts")
    if not isinstance(raw, dict):
        # Saved v0.4.x settings and direct callers remain valid.
        legacy_discipline = str(params.get("discipline") or "CIV").upper()
        legacy_count = int(params.get("count") or 0)
        raw = {legacy_discipline: legacy_count}

    counts: dict[str, int] = {}
    for discipline in SUPPORTED_DISCIPLINES:
        try:
            count = int(raw.get(discipline) or 0)
        except (TypeError, ValueError) as exc:
            raise ValueError(f"{discipline} RFI adedi gecersiz") from exc
        if count < 0 or count > MAX_RFI_COUNT:
            raise ValueError(
                f"{discipline} RFI adedi 0-{MAX_RFI_COUNT} arasinda olmali"
            )
        if count:
            counts[discipline] = count
    if not counts:
        raise ValueError("en az bir sheet icin RFI adedi secin")
    return counts


@register("rfi.log_add")
class AddRfiLogRows(Operation):
    title_key = "op.rfi.log_add.title"
    description_key = "op.rfi.log_add.desc"
    schema = ParamSchema(
        [
            FileParam(
                "workbook",
                "param.rfi.workbook",
                default=_default_workbook(),
                help_key="param.rfi.workbook.help",
            ),
            LongTextParam(
                "email_text",
                "param.rfi.emailText",
                default="",
                required=False,
                help_key="param.rfi.emailText.help",
            ),
            ChoiceParam(
                "email_provider",
                "param.rfi.emailProvider",
                choices=["offline", "gemini"],
                choice_labels={
                    "offline": "choice.rfi.emailProvider.offline",
                    "gemini": "choice.rfi.emailProvider.gemini",
                },
                default="offline",
                help_key="param.rfi.emailProvider.help",
            ),
            SecretParam(
                "gemini_api_key",
                "param.rfi.geminiApiKey",
                default="",
                required=False,
                visible_when={"email_provider": "gemini"},
                help_key="param.rfi.geminiApiKey.help",
            ),
            SheetCountsParam(
                "sheet_counts",
                "param.rfi.sheetCounts",
                choices=list(SUPPORTED_DISCIPLINES),
                default={"CIV": 3, "MEC": 0, "ELE": 0, "PRC": 0},
                help_key="param.rfi.sheetCounts.help",
            ),
            BoolParam(
                "overflow_next_day",
                "param.rfi.overflowNextDay",
                default=True,
                help_key="param.rfi.overflowNextDay.help",
            ),
        ]
    )

    def preview(self, params: dict) -> Plan:
        workbook_path = pathlib.Path(str(params["workbook"])).expanduser().resolve()
        email_text = str(params.get("email_text") or "").strip()
        wrap = bool(params.get("overflow_next_day", True))
        if workbook_path.suffix.lower() != ".xlsx" or not workbook_path.is_file():
            raise ValueError("gecerli bir .xlsx RFI Log dosyasi secin")
        counts = _requested_sheet_counts(params)
        detected = detect_discipline(email_text) if email_text else None
        has_new_selection = isinstance(params.get("sheet_counts"), dict)
        if detected in SUPPORTED_DISCIPLINES:
            # The default CIV=3 is a convenience, not an instruction. A pasted
            # MEC/ELE/PRC email should therefore route to its own sheet. Once
            # the user makes a genuine multi-sheet selection, keep every
            # selected sheet and add the email sheet if needed.
            if not has_new_selection or counts == {"CIV": 3}:
                counts = {detected: counts.get("CIV", 3)}
            elif detected not in counts:
                counts[detected] = 1

        requested = [
            discipline
            for discipline in SUPPORTED_DISCIPLINES
            if discipline in counts
        ]
        notify("progress", {"phase": "preview", "done": 0, "total": 4})
        sources = _read_sources(workbook_path, requested)
        notify("progress", {"phase": "preview", "done": 1, "total": 4})
        extraction: dict[str, Any] | None = None
        extraction_discipline: str | None = None
        extraction_notes: list[str] = []
        extraction_warnings: list[Warning] = []
        if email_text:
            extraction_discipline = (
                detected
                if detected in sources
                else requested[0]
            )
            assert extraction_discipline is not None
            source = sources[extraction_discipline]
            provider = str(params.get("email_provider") or "offline").lower()
            if provider not in {"offline", "gemini"}:
                provider = "offline"
            extraction, fallback_error = extract_email(
                email_text,
                provider=provider,
                api_key=str(params.get("gemini_api_key") or ""),
                discipline=extraction_discipline,
                itp_to_items=source["itpToItems"],
                known_tags=source["tags"],
                known_subcontractors=source["subcontractors"],
            )
            extracted_count = int(extraction.get("count") or 0)
            if 1 <= extracted_count <= MAX_RFI_COUNT:
                counts[extraction_discipline] = extracted_count
            raw_notes = extraction.get("notes")
            if isinstance(raw_notes, list):
                extraction_notes.extend(
                    str(note).strip() for note in raw_notes if str(note).strip()
                )
            confidence = float(extraction.get("confidence") or 0)
            used_provider = str(extraction.get("provider") or provider)
            extraction_warnings.append(
                Warning(
                    "warn.rfi.emailExtracted",
                    (
                        f"{used_provider} | {extraction_discipline} "
                        f"{counts[extraction_discipline]} RFI | "
                        f"confidence {confidence:.0%}"
                    ),
                )
            )
            if provider == "gemini":
                extraction_warnings.append(Warning("warn.rfi.geminiPrivacy"))
            if fallback_error:
                extraction_warnings.append(
                    Warning("warn.rfi.aiFallback", fallback_error)
                )

        start_day = _today() + timedelta(days=1)
        overflowed = 0
        issue_day = _today()
        rows: list[dict[str, Any]] = []
        batches: list[dict[str, Any]] = []
        for discipline in requested:
            source = sources[discipline]
            slots, sheet_overflowed = _allocate(
                source["occupied"],
                start_day,
                counts[discipline],
                wrap,
            )
            overflowed += sheet_overflowed
            sheet_rows: list[dict[str, Any]] = []
            base_rfis: list[str] = []
            for index in range(counts[discipline]):
                number = source["maxNumber"] + index + 1
                rfi = f"AZE-SUN-INT-QMS-RFI-{discipline}-{number:04d}"
                base_rfis.append(rfi)
                sheet_rows.append(
                    _record(
                        index=index,
                        discipline=discipline,
                        rfi=rfi,
                        source=source,
                        issue_day=issue_day,
                        slot=slots[index],
                    )
                )
            if extraction is not None and discipline == extraction_discipline:
                sheet_rows, validation_notes = _rows_from_email(
                    sheet_rows,
                    extraction,
                    source,
                )
                extraction_notes.extend(validation_notes)

            start_row = source["lastRow"] + 1
            if start_row + len(sheet_rows) - 1 > source["maxRow"]:
                raise ValueError(
                    f"{discipline} sheet'inde yeni satirlar icin "
                    "ayrilmis kapasite kalmadi"
                )
            batches.append(
                {
                    "discipline": discipline,
                    "startRow": start_row,
                    "records": copy.deepcopy(sheet_rows),
                    "baseRfis": base_rfis,
                    "presets": source["presets"],
                    "lookupColumns": source["lookupColumns"],
                    "maxRow": source["maxRow"],
                    "sourceData": {
                        key: source[key]
                        for key in (
                            "presets",
                            "itpToItems",
                            "lookupColumns",
                            "tags",
                            "areas",
                            "locations",
                            "subcontractors",
                            "maxRow",
                        )
                    },
                }
            )
            rows.extend(sheet_rows)

        if extraction_notes:
            extraction_warnings.append(
                Warning(
                    "warn.rfi.emailNeedsReview",
                    " | ".join(dict.fromkeys(extraction_notes))[:500],
                )
            )

        signature = _signature(workbook_path)
        backup_dir = workbook_path.parent / ".docflow-backups"
        backup_name = (
            f"{workbook_path.stem}_{datetime.now():%Y%m%d_%H%M%S}_"
            f"{uuid.uuid4().hex[:8]}{workbook_path.suffix}"
        )
        step_meta: dict[str, Any] = {
            "sourceSignature": signature,
            "batches": batches,
        }
        if len(batches) == 1:
            only = batches[0]
            step_meta.update(
                {
                    "discipline": only["discipline"],
                    "startRow": only["startRow"],
                    "records": only["records"],
                    "baseRfis": only["baseRfis"],
                    "presets": only["presets"],
                    "lookupColumns": only["lookupColumns"],
                    "sourceData": only["sourceData"],
                }
            )
        step = Step(
            "workbook_update",
            src=str(workbook_path),
            dst=str(backup_dir / backup_name),
            note_key="note.rfi.workbookUpdate",
            meta=step_meta,
        )
        meta = _plan_meta(sources, rows)
        copy_sources = {
            batch["discipline"]: _copy_source(
                workbook_path,
                batch["discipline"],
                int(batch["startRow"]),
                list(batch["records"]),
            )
            for batch in batches
        }
        meta["editor"]["copySources"] = copy_sources
        if len(copy_sources) == 1:
            meta["editor"]["copySource"] = next(iter(copy_sources.values()))
        if extraction is not None:
            meta["editor"]["emailExtraction"] = {
                "provider": extraction.get("provider"),
                "confidence": extraction.get("confidence"),
                "count": counts.get(extraction_discipline or "", 0),
                "sheet": extraction_discipline,
            }
        notify("progress", {"phase": "preview", "done": 4, "total": 4})
        return Plan(
            op_id=self.id,
            steps=[step],
            warnings=[
                *extraction_warnings,
                *_warnings_for_rows(rows, overflowed=overflowed, wrap=wrap),
            ],
            meta=meta,
        )

    def revise(self, plan: Plan, payload: dict) -> Plan:
        if len(plan.steps) != 1 or plan.steps[0].kind != "workbook_update":
            raise ValueError("gecersiz RFI Log plani")
        original = plan.steps[0]
        original_batches = list(original.meta.get("batches") or [])
        if not original_batches:
            raise ValueError("gecersiz coklu-sheet RFI plani")
        raw_rows = list(payload.get("rows") or [])
        allowed_sheets = {
            str(batch["discipline"]) for batch in original_batches
        }
        supplied_sheets = {
            str(row.get("_sheet") or "").upper()
            for row in raw_rows
            if isinstance(row, dict)
        }
        if supplied_sheets != allowed_sheets:
            raise ValueError("RFI sheet gruplari eklenemez veya silinemez")

        revised_batches: list[dict[str, Any]] = []
        revised_rows: list[dict[str, Any]] = []
        sources: dict[str, dict[str, Any]] = {}
        for batch in original_batches:
            discipline = str(batch["discipline"])
            source_data = copy.deepcopy(batch["sourceData"])
            sheet_rows = _safe_rows(
                [
                    row
                    for row in raw_rows
                    if isinstance(row, dict)
                    and str(row.get("_sheet") or "").upper() == discipline
                ],
                discipline=discipline,
                base_rfis=list(batch["baseRfis"]),
                source=source_data,
            )
            start_row = int(batch["startRow"])
            if start_row + len(sheet_rows) - 1 > int(batch["maxRow"]):
                raise ValueError(
                    f"{discipline} sheet'inde kopyalar icin kapasite kalmadi"
                )
            revised_batch = copy.deepcopy(batch)
            revised_batch["records"] = copy.deepcopy(sheet_rows)
            revised_batches.append(revised_batch)
            revised_rows.extend(sheet_rows)
            sources[discipline] = source_data

        if len(revised_rows) > 500:
            raise ValueError("tek planda en fazla 500 fiziksel satir olabilir")
        revised_step = copy.deepcopy(original)
        revised_step.meta["batches"] = revised_batches
        if len(revised_batches) == 1:
            revised_step.meta["records"] = copy.deepcopy(
                revised_batches[0]["records"]
            )
        meta = _plan_meta(sources, revised_rows)
        copy_sources = {
            batch["discipline"]: _copy_source(
                pathlib.Path(str(original.src)),
                str(batch["discipline"]),
                int(batch["startRow"]),
                list(batch["records"]),
            )
            for batch in revised_batches
        }
        meta["editor"]["copySources"] = copy_sources
        if len(copy_sources) == 1:
            meta["editor"]["copySource"] = next(iter(copy_sources.values()))
        original_extraction = (plan.meta.get("editor") or {}).get(
            "emailExtraction"
        )
        if original_extraction is not None:
            meta["editor"]["emailExtraction"] = copy.deepcopy(
                original_extraction
            )
        return Plan(
            op_id=self.id,
            steps=[revised_step],
            warnings=_warnings_for_rows(revised_rows),
            meta=meta,
        )

    def apply(self, plan: Plan) -> Result:
        if len(plan.steps) != 1 or plan.steps[0].kind != "workbook_update":
            return Result(
                failed=1,
                errors=[Warning("error.step_failed", "gecersiz plan")],
            )
        step = plan.steps[0]
        workbook_path = pathlib.Path(str(step.src))
        backup_path = pathlib.Path(str(step.dst))
        batches = list(step.meta.get("batches") or [])
        if not batches:
            return Result(
                failed=1,
                errors=[Warning("error.step_failed", "gecersiz sheet plani")],
            )
        signature = dict(step.meta["sourceSignature"])
        result = Result()
        if cancellation_requested():
            result.meta = {"cancelled": True, "pending": 1}
            result.errors.append(Warning("warn.operation.cancelled"))
            return result
        temp_path = workbook_path.with_name(
            f".{workbook_path.name}.{plan.plan_id[:8]}.docflow.tmp"
        )

        try:
            if _signature(workbook_path) != signature:
                raise RuntimeError(
                    "RFI Log onizlemeden sonra degismis. Yeni onizleme olusturun."
                )
            _assert_not_locked(workbook_path)
            backup_path.parent.mkdir(parents=True, exist_ok=True)
            if backup_path.exists():
                raise FileExistsError(str(backup_path))
            _copy_contents(workbook_path, backup_path)

            restore_step = Step(
                "restore_backup",
                src=str(backup_path),
                dst=str(workbook_path),
                note_key="note.rfi.restoreBackup",
                meta={"originalSha256": signature["sha256"]},
            )
            entry = JournalEntry(op_id=self.id, steps=[restore_step])
            result.journal_id = entry.journal_id
            write_journal(entry)

            notify("progress", {"phase": "apply", "done": 0, "total": 3})
            _write_batches(workbook_path, temp_path, batches)
            notify("progress", {"phase": "apply", "done": 1, "total": 3})
            _replace_with_retry(temp_path, workbook_path)
            for batch in batches:
                _verify_records(
                    workbook_path,
                    str(batch["discipline"]),
                    int(batch["startRow"]),
                    list(batch["records"]),
                )
            notify("progress", {"phase": "apply", "done": 2, "total": 3})

            restore_step.meta["appliedSha256"] = file_sha256(workbook_path)
            write_journal(entry)
            result.applied = sum(len(batch["records"]) for batch in batches)
            notify("progress", {"phase": "apply", "done": 3, "total": 3})
            return result
        except Exception as exc:  # noqa: BLE001
            temp_path.unlink(missing_ok=True)
            try:
                if (
                    backup_path.is_file()
                    and workbook_path.is_file()
                    and file_sha256(workbook_path) != signature["sha256"]
                ):
                    restore_temp = workbook_path.with_name(
                        f".{workbook_path.name}.docflow-restore.tmp"
                    )
                    _copy_contents(backup_path, restore_temp, overwrite=True)
                    _replace_with_retry(restore_temp, workbook_path)
            except Exception:
                pass
            result.failed = 1
            result.errors.append(_apply_error_warning(exc, workbook_path))
            return result
