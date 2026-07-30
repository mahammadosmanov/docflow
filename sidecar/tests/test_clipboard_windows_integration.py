"""Opt-in integration test for the real Excel COM clipboard renderer."""
from __future__ import annotations

import ctypes
import os
from ctypes import wintypes

import pytest
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from docflow_sidecar.clipboard import _copy_exact_from_excel
from docflow_sidecar.contracts import Plan


def _clipboard_text_and_formats() -> tuple[str, set[str]]:
    user32 = ctypes.WinDLL("user32", use_last_error=True)
    kernel32 = ctypes.WinDLL("kernel32", use_last_error=True)
    user32.OpenClipboard.argtypes = [wintypes.HWND]
    user32.OpenClipboard.restype = wintypes.BOOL
    user32.GetClipboardData.argtypes = [wintypes.UINT]
    user32.GetClipboardData.restype = wintypes.HANDLE
    user32.RegisterClipboardFormatW.argtypes = [wintypes.LPCWSTR]
    user32.RegisterClipboardFormatW.restype = wintypes.UINT
    user32.IsClipboardFormatAvailable.argtypes = [wintypes.UINT]
    user32.IsClipboardFormatAvailable.restype = wintypes.BOOL
    kernel32.GlobalLock.argtypes = [wintypes.HGLOBAL]
    kernel32.GlobalLock.restype = wintypes.LPVOID
    kernel32.GlobalUnlock.argtypes = [wintypes.HGLOBAL]

    assert user32.OpenClipboard(None)
    try:
        handle = user32.GetClipboardData(13)  # CF_UNICODETEXT
        assert handle
        pointer = kernel32.GlobalLock(handle)
        assert pointer
        try:
            text = ctypes.wstring_at(pointer)
        finally:
            kernel32.GlobalUnlock(handle)
        formats = {
            name
            for name in ("HTML Format", "Rich Text Format", "Biff12")
            if user32.IsClipboardFormatAvailable(
                user32.RegisterClipboardFormatW(name)
            )
        }
        return text, formats
    finally:
        user32.CloseClipboard()


@pytest.mark.skipif(
    os.name != "nt" or os.environ.get("DOCFLOW_RUN_EXCEL_INTEGRATION") != "1",
    reason="requires Microsoft Excel on an interactive Windows machine",
)
def test_excel_filters_exact_rfis_and_copies_header_with_native_format(tmp_path):
    path = tmp_path / "RFI Log.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "CIV"
    headers = [
        "RFI No.",
        "ITP No.",
        "Item No.",
        "INSPECTION DESCRIPTION",
        "FORMS",
        "TAG NO",
        "AREA",
        "DATE OF ISSUE",
        "INSPECTION DATE",
        "EXPIRY",
        "INSPECTION START TIME",
        "INSP. TYPE AZFEN",
        "INSP. TYPE INTEC",
        "Inspection Location",
        "SUBCONTRACTOR",
    ]
    for column, value in enumerate(headers, 1):
        cell = sheet.cell(9, column, value)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4F4F4F")
        sheet.column_dimensions[cell.column_letter].width = 16
    rows = {
        10: ("AZE-SUN-INT-QMS-RFI-CIV-0100", "OLD"),
        12: ("AZE-SUN-INT-QMS-RFI-CIV-0861", "SECOND"),
        14: ("AZE-SUN-INT-QMS-RFI-CIV-0860", "FIRST"),
    }
    for row_number, (rfi, form) in rows.items():
        sheet.cell(row_number, 1, rfi)
        sheet.cell(row_number, 5, form)
        sheet.cell(row_number, 14, "North Site")
        sheet.cell(row_number, 15, "AZFEN")
    workbook.save(path)
    workbook.close()

    plan = Plan(
        op_id="rfi.log_add",
        meta={
            "editor": {
                "kind": "recordTable",
                "copyEnabled": True,
                "copySource": {
                    "workbook": str(path),
                    "sheet": "CIV",
                    "headerRow": 9,
                    "startRow": 12,
                    "rfiNumbers": [
                        "AZE-SUN-INT-QMS-RFI-CIV-0860",
                        "AZE-SUN-INT-QMS-RFI-CIV-0861",
                    ],
                    "columns": [
                        1, 2, 3, 4, 5, 6, 7, 8, 9, 11, 12, 13, 14, 15
                    ],
                },
            }
        },
    )

    # Excel stays open intentionally: that is the user-facing preview and it
    # keeps Excel's complete native clipboard formats alive for Outlook.
    assert _copy_exact_from_excel(plan, 2) is True
    clipboard_text, clipboard_formats = _clipboard_text_and_formats()
    assert "RFI No." in clipboard_text
    assert "SUBCONTRACTOR" in clipboard_text
    assert "AZE-SUN-INT-QMS-RFI-CIV-0860" in clipboard_text
    assert "AZE-SUN-INT-QMS-RFI-CIV-0861" in clipboard_text
    assert "AZE-SUN-INT-QMS-RFI-CIV-0100" not in clipboard_text
    assert "EXPIRY" not in clipboard_text
    assert clipboard_formats == {"HTML Format", "Rich Text Format", "Biff12"}
